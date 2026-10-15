"""Export only public presentation fields, never arbitrary report files."""
from io import BytesIO


COLUMNS = (
    ("code", "代码"), ("name", "股票"), ("source_layer", "原最终层级"),
    ("board_scope", "交易范围"), ("close", "快照收盘"), ("pct_chg", "涨跌幅(%)"),
    ("primary_direction", "方向"), ("observe_low", "观察下限"),
    ("observe_high", "观察上限"), ("invalid_price", "失效价"),
    ("pressure_price", "压力位"), ("entry_reason", "入选依据"),
    ("risk_reasons", "风险"), ("price_reference_date", "价格参考日期"),
)


def cell(value):
    if isinstance(value, list):
        return "；".join(str(item) for item in value)
    return "—" if value is None else str(value)


def md(value):
    return cell(value).replace("|", " / ").replace("\n", "；").replace("\r", "")


def markdown(data):
    lines = [f"# 收盘复盘公开明细 · {data['trade_date']}", "",
             "来源：网页公开数据。缺失项不补写，不替代原始审计附件或实际交易记录。", "",
             "## 市场", "", f"- 数据范围：{data['market_scope']}",
             f"- 状态：{md(data['market'].get('status'))}",
             f"- 评分：{md(data['market'].get('score'))}",
             f"- 上涨/下跌：{md(data['market'].get('up_count'))}/{md(data['market'].get('down_count'))}",
             f"- 允许新增推荐：{'是' if data['new_recommendations_allowed'] else '否'}", "",
             "## 完整当日记录", "", "包含观察、回避及研究记录，不等于推荐名单。", "",
             "| " + " | ".join(label for _, label in COLUMNS) + " |",
             "| " + " | ".join("---" for _ in COLUMNS) + " |"]
    for row in data["candidates"]:
        lines.append("| " + " | ".join(md(row.get(key)) for key, _ in COLUMNS) + " |")
    lines.extend(["", "## 限制", "", "历史价格有效性及复权口径须单独核验。事件检索不代表因果解释。"])
    return "\n".join(lines).encode("utf-8")


def workbook(data):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    book = Workbook()
    sheet = book.active
    sheet.title = "当日公开记录"
    sheet.append([label for _, label in COLUMNS])
    for item in data["candidates"]:
        sheet.append([cell(item.get(key)) if isinstance(item.get(key), list)
                      else item.get(key) for key, _ in COLUMNS])
    # Text must remain literal even when source strings begin with '='.
    for row in sheet.iter_rows():
        for value in row:
            if isinstance(value.value, str):
                value.data_type = "s"
            value.alignment = Alignment(vertical="top", wrap_text=True)
    for value in sheet[1]:
        value.fill = PatternFill("solid", fgColor="176650")
        value.font = Font(color="FFFFFF", bold=True)
    for index, (key, _) in enumerate(COLUMNS, 1):
        sheet.column_dimensions[get_column_letter(index)].width = 48 if key in {"entry_reason", "risk_reasons"} else 18
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    notes = book.create_sheet("口径说明")
    notes.append(["复盘日期", data["trade_date"]])
    notes.append(["数据范围", data["market_scope"]])
    notes.append(["口径", "完整当日公开候选记录，包含观察及回避，不等于推荐名单。"])
    notes.append(["限制", "价格引用原快照日期，跨期有效性未核验；不是原板块趋势Excel。"])
    notes.column_dimensions["A"].width = 20
    notes.column_dimensions["B"].width = 90
    stream = BytesIO()
    book.save(stream)
    return stream.getvalue()
